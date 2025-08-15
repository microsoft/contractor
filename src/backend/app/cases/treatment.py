import os
import json
from docx import Document

PASTA = './transliteracoes'

# Carregar avaliações
with open(os.path.join(PASTA, 'evaluations.json'), encoding='utf-8') as f:
    avaliacoes_json = json.load(f)

avaliacoes_map = {}
for bloco in avaliacoes_json:
    for av in bloco.get('avaliacoes', []):
        avaliacoes_map[av['id'].strip().lower()] = av

# Listar arquivos de transliteração
arquivos = [arq for arq in os.listdir(PASTA) if arq.startswith('transliteration_') and arq.endswith('.json')]

# Carregar transliterações
translit = []
for arq in arquivos:
    with open(os.path.join(PASTA, arq), encoding='utf-8') as f:
        data = json.load(f)
        if isinstance(data, dict) and 'transliteration' in data:
            translit.extend(data['transliteration'])

# Criar documento Word
doc = Document()
doc.add_heading('Consolidação de Traduções e Avaliações', 0)

for frase in translit:
    frase_id = str(frase.get('id', '')).strip()
    frase_port = frase.get('portugues', '').strip()
    doc.add_heading(f'Frase {frase_id}: {frase_port}', level=1)
    for idioma in ['katukina', 'kulina', 'marubo', 'matses', 'pano', 'matis', 'korubo', 'kanamari']:
        if idioma in frase:
            doc.add_heading(idioma.capitalize(), level=2)
            doc.add_paragraph(f'Tradução: {frase[idioma]}')
            # Procurar avaliação
            chaves = [
                f'{frase_id}-{idioma}', f'{frase_id} - {idioma}', f'{frase_id}.{idioma}',
                f'{idioma}', frase_id
            ]
            avaliacao = None
            for chave in chaves:
                if chave.lower() in avaliacoes_map:
                    avaliacao = avaliacoes_map[chave.lower()]
                    break
            if avaliacao:
                doc.add_paragraph(f"Avaliação: {avaliacao.get('comentario','')}")
                doc.add_paragraph(f"Status: {avaliacao.get('status','')}")
            else:
                doc.add_paragraph("Avaliação: Não encontrada.")
                doc.add_paragraph("Status: Desconhecido.")

# Salvar
doc.save('consolidado_traducao_avaliacao.docx')
print('Documento gerado: consolidado_traducao_avaliacao.docx')
